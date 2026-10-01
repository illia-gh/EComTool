"""Importable memory-first boundary around legacy Stage-1 calculations.

The legacy script still owns formulas/profile construction. This module runs it
without XLSX serialization and converts its canonical output table directly to
immutable Stage-2 inputs. Legacy CLI/MATLAB execution remains unchanged.
"""
from __future__ import annotations

import io
import os
import runpy
import threading
from contextlib import redirect_stdout
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

import pandas as pd

from ecom_port.frontend import Inputs, inputs_from_frame


LogCallback = Callable[[str], None]
_EXECUTION_LOCK = threading.Lock()


@dataclass(frozen=True)
class PreparedData:
    """Immutable calculation snapshot produced before workbook serialization."""

    fingerprint: str
    inputs: Inputs
    legacy_frame: pd.DataFrame | None = None

    def compact(self) -> "PreparedData":
        """Release large legacy export table while retaining Stage-2 arrays."""
        if self.legacy_frame is None:
            return self
        return PreparedData(self.fingerprint, self.inputs)


def prepare_data(input_path: Path, stage1_dir: Path, fingerprint: str, *,
                 keep_export_frame: bool = False,
                 on_log: LogCallback | None = None) -> PreparedData:
    """Run unchanged Stage-1 formulas and return arrays before XLSX writing."""
    input_path = Path(input_path).resolve()
    stage1_dir = Path(stage1_dir).resolve()
    script_path = Path(__file__).resolve().with_name("ecomtool_stage_1.py")
    capture = io.StringIO()
    namespace: dict | None = None

    with _EXECUTION_LOCK:
        previous_cwd = Path.cwd()
        previous_input = os.environ.get("ECOM_STAGE1_INPUT")
        previous_output = os.environ.get("ECOM_STAGE1_OUTPUT")
        previous_skip = os.environ.get("ECOM_STAGE1_SKIP_XLSX")
        try:
            os.environ["ECOM_STAGE1_INPUT"] = str(input_path)
            os.environ["ECOM_STAGE1_OUTPUT"] = str(stage1_dir / "EComTool_Output.xlsx")
            os.environ["ECOM_STAGE1_SKIP_XLSX"] = "1"
            os.chdir(stage1_dir)
            with redirect_stdout(capture):
                namespace = runpy.run_path(str(script_path), run_name="_ecom_stage1_core")
        finally:
            os.chdir(previous_cwd)
            for name, value in (
                ("ECOM_STAGE1_INPUT", previous_input),
                ("ECOM_STAGE1_OUTPUT", previous_output),
                ("ECOM_STAGE1_SKIP_XLSX", previous_skip),
            ):
                if value is None:
                    os.environ.pop(name, None)
                else:
                    os.environ[name] = value
            if on_log is not None and capture.getvalue():
                on_log(capture.getvalue())

    if namespace is None or not isinstance(namespace.get("df_output"), pd.DataFrame):
        raise RuntimeError("Stage-1 completed without producing canonical output table")
    frame = namespace["df_output"]
    inputs = inputs_from_frame(frame, normalize_excel=True)
    return PreparedData(
        fingerprint=fingerprint,
        inputs=inputs,
        legacy_frame=frame if keep_export_frame else None,
    )


def export_stage1_xlsx(data: PreparedData, output_path: Path) -> None:
    """Write optional legacy workbook from captured Stage-1 snapshot."""
    if data.legacy_frame is None:
        raise ValueError("PreparedData does not retain legacy export table")
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    # Keep writer aligned with legacy workbook layout and formatting contract.
    import openpyxl
    import numpy as np
    from openpyxl.writer.excel import ExcelWriter
    from zipfile import ZIP_DEFLATED, ZipFile

    workbook = openpyxl.Workbook(write_only=True)
    worksheet = workbook.create_sheet(title="Sheet1")
    internal_headers = {"Other Input Data", "Run Metadata"}
    for column_number, header in enumerate(data.legacy_frame.iloc[0], start=1):
        if str(header).strip() in internal_headers:
            column_letter = openpyxl.utils.get_column_letter(column_number)
            worksheet.column_dimensions[column_letter].hidden = True
    for row in data.legacy_frame.itertuples(index=False, name=None):
        worksheet.append([
            None if isinstance(value, (float, np.floating)) and np.isnan(value) else value
            for value in row
        ])
    archive = ZipFile(output_path, "w", ZIP_DEFLATED, allowZip64=True, compresslevel=1)
    try:
        ExcelWriter(workbook, archive).save()
    finally:
        archive.close()

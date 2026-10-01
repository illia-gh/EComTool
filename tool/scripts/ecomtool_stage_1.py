import pandas as pd
import json
import os
import sys
import time
import warnings
import numpy as np
import openpyxl
from openpyxl.writer.excel import ExcelWriter
from zipfile import ZIP_DEFLATED, ZipFile

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
if SCRIPT_DIR not in sys.path:
    sys.path.insert(0, SCRIPT_DIR)
from tariff_io import read_tariff_profile, require_positive_number


def load_green_tariff_rate():
    """Read canonical Green Tariff rate from shipped Stage-2 config."""
    config_path = os.path.normpath(
        os.path.join(SCRIPT_DIR, "..", "defaults", "sc_default.json")
    )
    with open(config_path, "r", encoding="utf-8") as config_file:
        value = json.load(config_file)["green_tariff_EUR_per_kWh"]
    rate = float(value)
    if not np.isfinite(rate) or rate <= 0:
        raise ValueError("green_tariff_EUR_per_kWh must be finite and greater than zero")
    return rate

stage1_started = time.perf_counter()

warnings.filterwarnings(
    "ignore",
    message="Data Validation extension is not supported and will be removed",
)

# --- Input Data and Setup ---

input_excel = os.environ.get("ECOM_STAGE1_INPUT", "EComTool_User_Input.xlsx")
output_excel = os.environ.get("ECOM_STAGE1_OUTPUT", "EComTool_Output.xlsx")

df_input = pd.read_excel(input_excel, header=None)

# Read data from specific cells
country = str(df_input.iloc[3, 0]) # cell A4
full_city_name = str(df_input.iloc[3, 1]) # cell B4
area_type = str(df_input.iloc[3, 2]) # cell C4
tariff_from = str(df_input.iloc[3, 3]) # cell D4
tariff_to = str(df_input.iloc[3, 4]) # cell E4
tariff_fixed_from_ua = 7.50 / 52.0 # EUR/kWh, all-in Ukrainian fixed import tariff
tariff_fixed_from_lv = 0.15 # EUR/kWh
tariff_fixed_to_ua = 0.05 # EUR/kWh
tariff_fixed_to_lv = 0.075 # EUR/kWh
tariff_green_ua = load_green_tariff_rate() # EUR/kWh; canonical source is tool/defaults/sc_default.json
tariff_fixed_from_own = float(df_input.iloc[90, 1]) if pd.notna(df_input.iloc[90, 1]) else 0.0 # EUR/kWh, cell B91
tariff_fixed_to_own = require_positive_number(df_input.iloc[91, 1], "B92") # EUR/kWh, cell B92

pv_p_existing = df_input.iloc[9, 1] # cell B10
pv_p_projected = df_input.iloc[10, 1] # cell B11
pv_adoption_traj = df_input.iloc[96, 1] # cell B97
pv_prof = str(df_input.iloc[11, 1]) # cell B12
pv_age = df_input.iloc[92, 1] # cell B93
pv_degradation = df_input.iloc[93, 1] # cell B94

ev_num_existing = df_input.iloc[13, 1] # cell B14
ev_num_projected = df_input.iloc[14, 1] # cell B15
ev_adoption_traj = df_input.iloc[97, 1] # cell B98
ev_home_prof = str(df_input.iloc[15, 1]) # cell B16
ev_public_prof = str(df_input.iloc[16, 1]) # cell B17
ev_work_prof = str(df_input.iloc[17, 1]) # cell B18
ev_fast_prof = str(df_input.iloc[18, 1]) # cell B19
ev_home_share = df_input.iloc[19, 1] # cell B20
ev_public_share = df_input.iloc[20, 1] # cell B21
ev_work_share = df_input.iloc[21, 1] # cell B22
ev_fast_share = df_input.iloc[22, 1] # cell B23
ev_smart_share = df_input.iloc[23, 1] # cell B24
smart_ev_home_prof = str(df_input.iloc[24, 1]) # cell B25
smart_ev_public_prof = str(df_input.iloc[25, 1]) # cell B26

expected_bess_capacity_labels = {
    27: "Existing BESS Installed Capacity, kWh",
    28: "Projected BESS Installed Capacity, kWh",
}
for row_index, expected_label in expected_bess_capacity_labels.items():
    actual_label = str(df_input.iloc[row_index, 0]).strip()
    if actual_label != expected_label:
        raise ValueError(
            f"Input!A{row_index + 1} must be {expected_label!r}; got {actual_label!r}"
        )
bess_capacity_existing_kWh = df_input.iloc[27, 1] # cell B28
bess_capacity_projected_kWh = df_input.iloc[28, 1] # cell B29
bess_degradation_pct = float(df_input.iloc[94, 1]) # cell B95
if not np.isfinite(bess_degradation_pct) or not 0 <= bess_degradation_pct < 100:
    raise ValueError("B95 BESS degradation must be finite and in [0, 100) percent")
bess_degradation_rate = bess_degradation_pct / 100.0

planning_horizon = int(df_input.iloc[95, 1]) # cell B96
annual_load_growth_rate = float(df_input.iloc[98, 1]) if pd.notna(df_input.iloc[98, 1]) else 0.0 # cell B99

# Get headers for the output file
header_1 = str(df_input.iloc[11, 0]) # cell A12
header_2 = str(df_input.iloc[15, 0]) # cell A16
header_3 = str(df_input.iloc[16, 0]) # cell A17
header_4 = str(df_input.iloc[17, 0]) # cell A18
header_5 = str(df_input.iloc[18, 0]) # cell A19
header_6 = str(df_input.iloc[24, 0]) # cell A25
header_7 = str(df_input.iloc[25, 0]) # cell A26

# Extract the first word from the city name
city = full_city_name.split()[0]

print(f"Country: {country}, City: {city}, Profile Source: {pv_prof}")

# --- STEP 1: Determine the PV folder path ---

path_to_pv_folder = ""

if "Database" in pv_prof:
    if country == "Ukraine":
        path_to_pv_folder = os.path.join("Database", "PV Profiles", "UA_PV_generation") # os.path allows to use the script on macOS/Linux
    elif country == "Latvia":
        path_to_pv_folder = os.path.join("Database", "PV Profiles", "LV_PV_generation")
elif "Own" in pv_prof:
    path_to_pv_folder = os.path.join("User Own Input", "PV Profiles")

# --- STEP 2: File search and data reading ---

l_pv_data = [] # final generation profile
target_pv_file = ""
skip_rows_count = 0
unit_converter = 1 # default converter (no changes for kW files)

# Define which file to look for based on pv_prof
if "PVGIS" in pv_prof:
    l_pv_files = os.listdir(path_to_pv_folder)
    for pv_file_name in l_pv_files:
        if "Database" in pv_prof:
            if city in pv_file_name and "ninja" not in pv_file_name:
                target_pv_file = pv_file_name
                break
        elif "Own" in pv_prof:
            if "own_pvgis" in pv_file_name.lower():
                target_pv_file = pv_file_name
                break
    skip_rows_count = 10
    start_row, end_row, col_idx = 1, 8761, 1
    unit_converter = 1000 # convert W to kW for PVGIS files

elif "Renewables.Ninja" in pv_prof:
    l_pv_files = os.listdir(path_to_pv_folder)
    for pv_file_name in l_pv_files:
        if "Database" in pv_prof:
            if city in pv_file_name and "ninja" in pv_file_name:
                target_pv_file = pv_file_name
                break
        elif "Own" in pv_prof:
            if "own_ninja" in pv_file_name.lower():
                target_pv_file = pv_file_name
                break
    skip_rows_count = 3
    start_row, end_row, col_idx = 1, 8761, 2
    unit_converter = 1 # Ninja is already in kW

elif "standard format, 1 hour" in pv_prof:
    target_pv_file = "own_PV_profile_1_hour.csv"
    skip_rows_count = 0
    start_row, end_row, col_idx = 1, 8761, 0

elif "standard format, 15 min" in pv_prof:
    target_pv_file = "own_PV_profile_15_min.csv"
    skip_rows_count = 0
    start_row, end_row, col_idx = 1, 35041, 0

# Check if the file was found
if target_pv_file != "" and os.path.exists(os.path.join(path_to_pv_folder, target_pv_file)):
    full_pv_path = os.path.join(path_to_pv_folder, target_pv_file)
    print(f"Processing file: {target_pv_file}")

    df_pv_csv = pd.read_csv(full_pv_path, header=None, sep=None, engine='python', skiprows=skip_rows_count)

    # Apply unit_converter during list creation
    l_raw_pv_values = []
    pv_values = df_pv_csv.iloc[start_row:end_row, col_idx].tolist()
    for pv_value in pv_values:
        converted_pv_value = float(pv_value) / unit_converter
        l_raw_pv_values.append(converted_pv_value)

    # --- STEP 3: PV Data Processing and Scaling ---

    if "15 min" in pv_prof:
        l_pv_data = l_raw_pv_values
    else:
        for pv_value in l_raw_pv_values:
            l_pv_data.extend([pv_value] * 4)
else:
    print(f"CRITICAL ERROR: PV file {target_pv_file} not found in {path_to_pv_folder}")
    sys.exit(1)

if l_pv_data:
    # Calculate the scaled PV generation profile (beginning of Year 1 - this is a baseline)
    pv_p_existing_scaled = pv_p_existing * (1 - pv_degradation/100) ** pv_age
    l_pv_data_scaled = []
    for pv_value in l_pv_data:
        l_pv_data_scaled.append(pv_value * pv_p_existing_scaled)

# --- STEP 4: EV Profiles Processing ---

d_ev_profiles_results = {}
# List of tuples of profile types for processing
l_ev_types = [
    ('home', ev_home_prof, "regular_home"),
    ('public', ev_public_prof, "regular_public"),
    ('work', ev_work_prof, "regular_workplace"),
    ('fast', ev_fast_prof, "regular_fast"),
    ('smart_home', smart_ev_home_prof, "smart_home"),
    ('smart_public', smart_ev_public_prof, "smart_public")
]

# Process each EV profile one by one from the list above
for t_ev_tuple in l_ev_types:
    ev_profile_key = t_ev_tuple[0] # internal ID used to store the data in the d_ev_profiles_results
    ev_user_choice = t_ev_tuple[1] # user selection in Excel
    ev_filename_part = t_ev_tuple[2] # descriptive name part to identify the existing CSV file

    l_ev_data = []
    if "Database" in ev_user_choice or "Own" in ev_user_choice:
        # Determine path
        if "Database" in ev_user_choice:
            path_to_ev_folder = os.path.join("Database", "EV Charging Profiles")
            prefix = "NL_charging_profile_"
        else:
            path_to_ev_folder = os.path.join("User Own Input", "EV Charging Profiles")
            prefix = "own_charging_profile_"

        # Determine filename
        if "15 min" in ev_user_choice or "Database" in ev_user_choice:
            time_suffix = "_15_min"
        else:
            time_suffix = "_1_hour"

        target_ev_file = f"{prefix}{ev_filename_part}{time_suffix}.csv"

        # Check for file existence (lower case search)
        found_ev_file = ""
        if os.path.exists(path_to_ev_folder):
            for f in os.listdir(path_to_ev_folder):
                if target_ev_file.lower() == f.lower():
                    found_ev_file = f
                    break

        if found_ev_file:
            df_ev_csv = pd.read_csv(os.path.join(path_to_ev_folder, found_ev_file), header=None, skiprows=1)
            raw_ev_values = df_ev_csv.iloc[:, 0].tolist()

            # Convert to 15-min format if necessary
            if "1 hour" in ev_user_choice:
                for ev_value in raw_ev_values:
                    l_ev_data.extend([float(ev_value)] * 4)
            else:
                for v in raw_ev_values:
                    l_ev_data.append(float(v)) # convert each value to a floating-point number
        else:
            print(f"CRITICAL ERROR: EV file {target_ev_file} not found in {path_to_ev_folder}")
            sys.exit(1)
    else:
        # If profile is not selected, fill with zeros
        l_ev_data = [0.0] * 35040

    d_ev_profiles_results[ev_profile_key] = l_ev_data

# --- STEP 5: Aggregated EV Profile and Scaling ---

l_ev_aggregated_shape = []
smart_ratio = ev_smart_share / 100

for i in range(35040):
    # Calculation based on the weighted aggregation formula (normalized to 1 EV)
    ev_reg_part = (d_ev_profiles_results['home'][i] * ev_home_share + d_ev_profiles_results['public'][i] * ev_public_share) * (1 - smart_ratio)
    ev_smart_part = (d_ev_profiles_results['smart_home'][i] * ev_home_share + d_ev_profiles_results['smart_public'][i] * ev_public_share) * smart_ratio
    ev_other_part = d_ev_profiles_results['work'][i] * ev_work_share + d_ev_profiles_results['fast'][i] * ev_fast_share

    ev_value_unit = (ev_reg_part + ev_smart_part + ev_other_part) / 100
    l_ev_aggregated_shape.append(ev_value_unit)

l_ev_data_aggregated = []
for v in l_ev_aggregated_shape:
    value = v * ev_num_existing
    l_ev_data_aggregated.append(value)

# --- STEP 6: Participants' Load Profiles Processing ---

l_load_configs = [] # list of dictionaries for information about ECom participants
l_load_data = [] # list to store the 15-min profiles for the output saving

# Mapping load type strings to database sheet names
d_load_type_mapping = {
    "Private House (medium small, < 10,000 kWh/annum)": {"Latvia": ["LV_private_h._2", "LV_private_h._7", "LV_private_h._8"], "Ukraine": ["UA_private_h._1"]},
    "Private House (medium big, 10,000-20,000 kWh/annum)": {"Latvia": ["LV_private_h._1", "LV_private_h._4", "LV_private_h._5"], "Ukraine": ["UA_private_h._1"]},
    "Private House (big, >20,000 kWh/annum)": {"Latvia": ["LV_private_h._3", "LV_private_h._6"], "Ukraine": ["UA_private_h._1"]},
    "Apartment Building (16 apt, low energy efficiency)": {"Latvia": ["LV_apt_16_2"], "Ukraine": ["LV_apt_16_2"]},
    "Apartment Building (16 apt, high energy efficiency)": {"Latvia": ["LV_apt_16_1"], "Ukraine": ["LV_apt_16_1"]},
    "Apartment Building (27 apt, low energy efficiency)": {"Latvia": ["LV_apt_27_2"], "Ukraine": ["LV_apt_27_2"]},
    "Apartment Building (27 apt, high energy efficiency)": {"Latvia": ["LV_apt_27_1"], "Ukraine": ["LV_apt_27_1"]},
    "Apartment Building (31 apt, low energy efficiency)": {"Latvia": ["LV_apt_31_2"], "Ukraine": ["LV_apt_31_2"]},
    "Apartment Building (31 apt, high energy efficiency)": {"Latvia": ["LV_apt_31_1"], "Ukraine": ["LV_apt_31_1"]},
    "Apartment Building (54 apt, low energy efficiency)": {"Latvia": ["LV_apt_54_2"], "Ukraine": ["LV_apt_54_2"]},
    "Apartment Building (54 apt, high energy efficiency)": {"Latvia": ["LV_apt_54_1"], "Ukraine": ["LV_apt_54_1"]},
    "Apartment Building (66 apt, low energy efficiency)": {"Latvia": ["LV_apt_66_2"], "Ukraine": ["LV_apt_66_2"]},
    "Apartment Building (66 apt, high energy efficiency)": {"Latvia": ["LV_apt_66_1"], "Ukraine": ["LV_apt_66_1"]},
    "Apartment Building (80 apt, low energy efficiency)": {"Latvia": ["LV_apt_80_1"], "Ukraine": ["LV_apt_80_1"]},
    "Shop (small, < 150,000 kWh/annum)": {"Latvia": ["LV_shop_1"], "Ukraine": ["LV_shop_1"]},
    "Shop (big, > 200,000 kWh/annum)": {"Latvia": ["LV_shop_2"], "Ukraine": ["LV_shop_2"]},
    "Hotel (≈ 100,000 kWh/annum)": {"Latvia": ["LV_hotel"], "Ukraine": ["LV_hotel"]},
    "Bank (≈ 100,000 kWh/annum)": {"Latvia": ["LV_bank"], "Ukraine": ["LV_bank"]},
    "School (≈ 55,000 kWh/annum)": {"Latvia": ["LV_school"], "Ukraine": ["LV_school"]},
    "Kindergarten (≈ 20,000 kWh/annum)": {"Latvia": ["LV_kindergarten"], "Ukraine": ["LV_kindergarten"]},
    "Industry (small, ≈ 100,000 kWh/annum)": {"Latvia": ["LV_industry"], "Ukraine": ["LV_industry"]}
}

l_load_data_aggregated_base = [0.0] * 35040

# Range 35-84 corresponds to row indices 34 to 83 in a 0-indexed dataframe
for row_idx in range(34, 84):
    load_type = df_input.iloc[row_idx, 2] # column C (participant's type)

    # Check if the row contains valid load data (skip if the column C is empty)
    if pd.notna(load_type) and str(load_type).strip() != "":
        d_load_entry = {
            'load_id': str(df_input.iloc[row_idx, 0]),                                                             # column A (index 0)
            'load_type': str(load_type),
            'load_num': int(df_input.iloc[row_idx, 3]) if pd.notna(df_input.iloc[row_idx, 3]) else 1,              # column D (index 3)
            'load_annual_cons': float(df_input.iloc[row_idx, 4]) if pd.notna(df_input.iloc[row_idx, 4]) else None, # column E (index 4), handling NaN values
            'load_prof': str(df_input.iloc[row_idx, 5])                                                            # column F (index 5)
        }
        l_load_configs.append(d_load_entry)

        # --- Profile Extraction and Processing ---
        l_processed_load_prof = []

        if "Database" in d_load_entry['load_prof']:
            db_path = os.path.join("Database", "Load Profiles", "standard_load_profiles_database.xlsx")
            if os.path.exists(db_path):
                db_target_sheets = d_load_type_mapping.get(d_load_entry['load_type'], {}).get(country, [])
                l_temp_profiles = []

                for sheet in db_target_sheets:
                    df_sheet = pd.read_excel(db_path, sheet_name=sheet, header=None)

                    if len(df_sheet) >= 8788:
                        # LEAP YEAR detected: exclude Feb 29 (rows 1421-1444 in Excel)
                        val_last = df_sheet.iloc[8787, 1]
                        # Combine main parts while skipping the leap day slice
                        vals_before_leap = df_sheet.iloc[4:1420, 1].tolist()
                        vals_after_leap = df_sheet.iloc[1444:8787, 1].tolist()
                        l_combined_load_prof = [val_last] + vals_before_leap + vals_after_leap
                    else:
                        # NON-LEAP YEAR detected (8760 data points + headers)
                        val_last = df_sheet.iloc[8763, 1]
                        vals_main = df_sheet.iloc[4:8763, 1].tolist()
                        l_combined_load_prof = [val_last] + vals_main

                    # Scaling multiplier
                    if d_load_entry['load_annual_cons'] is None:
                        load_multiplier = d_load_entry['load_num']
                    else:
                        db_standard_annual_cons = float(df_sheet.iloc[2, 1])
                        load_multiplier = (d_load_entry['load_annual_cons'] * d_load_entry['load_num']) / db_standard_annual_cons

                    l_load_data_scaled = []
                    for load_value in l_combined_load_prof:
                        scaled_load_value = float(load_value) * load_multiplier
                        l_load_data_scaled.append(scaled_load_value)

                    l_temp_profiles.append(l_load_data_scaled)

                if l_temp_profiles:
                    num_sheets = len(l_temp_profiles)
                    # Create a single representative load profile for a category when the db provides multiple sample profiles (i.e., for private houses)
                    l_hourly_avg = []
                    # "Transpose" the list of lists so each 'x' contains the values for the same hour across all sheets
                    transposed_profiles = zip(*l_temp_profiles) # each sub-list is an 8760-hour profile
                    for x in transposed_profiles:
                        load_hourly_avg = sum(x) / num_sheets
                        l_hourly_avg.append(load_hourly_avg)

                    for h_val in l_hourly_avg:
                        l_processed_load_prof.extend([h_val] * 4)

        elif "Own" in d_load_entry['load_prof']:
            own_load_folder = os.path.join("User Own Input", "Load Profiles")
            formatted_id = d_load_entry['load_id'].lower().replace(" ", "_")

            if "15 min" in d_load_entry['load_prof']:
                target_load_file = f"own_{formatted_id}_15_min.csv"
                load_prof_length=35041
            elif "1 hour" in d_load_entry['load_prof']:
                target_load_file = f"own_{formatted_id}_1_hour.csv"
                load_prof_length=8761
            else:
                print(f"CRITICAL ERROR: unsupported own load profile format: {d_load_entry['load_prof']}")
                sys.exit(1)

            full_load_path = os.path.join(own_load_folder, target_load_file)
            if os.path.exists(full_load_path):
                df_own = pd.read_csv(full_load_path, header=None)
                l_own_raw_load_prof = df_own.iloc[1:load_prof_length, 0].astype(float).tolist()

                if d_load_entry['load_annual_cons'] is None:
                    load_multiplier = d_load_entry['load_num']
                else:
                    # Calculate current annual consumption from the own profile
                    own_annual_cons = sum(l_own_raw_load_prof)
                    if "15 min" in d_load_entry['load_prof']:
                        own_annual_cons = own_annual_cons / 4 # Sum of power values to energy (kWh)

                    # Check for zeros (if a user accidentally provides a file full of zeros)
                    if own_annual_cons != 0:
                        load_multiplier = d_load_entry['load_annual_cons'] * d_load_entry['load_num'] / own_annual_cons
                    else:
                        load_multiplier = d_load_entry['load_num']

                scaled_load_values = []
                for v in l_own_raw_load_prof:
                    scaled_load_values.append(v * load_multiplier)

                if "1 hour" in d_load_entry['load_prof']:
                    for v in scaled_load_values:
                        l_processed_load_prof.extend([v] * 4)
                else:
                    l_processed_load_prof = scaled_load_values

        # Fallback to zeros if profile processing failed
        if not l_processed_load_prof:
            l_processed_load_prof = [0.0] * 35040

        # Add the current participant's profile to the aggregated list
        for i in range(35040):
            l_load_data_aggregated_base[i] += l_processed_load_prof[i]

        l_load_data.append({
            'load_info_header': [d_load_entry['load_id'], d_load_entry['load_prof']],
            'load_info_data': l_processed_load_prof
        })

print(f"Detected {len(l_load_configs)} rows of input about ECom participants.")
print(l_load_configs)
print(f"Processed {len(l_load_data)} participant profiles.")

# --- STEP 7: Final Saving with Adoption Trajectories ---

# Adoption functions
def calculate_pv_adoption(t, existing, projected, planning_horizon, traj_type):
    if traj_type == "Instantaneous":
        return existing + projected
    elif traj_type == "Linear":
        return existing + (projected / planning_horizon) * t
    elif traj_type == "Exponential":
        if existing <= 0:
            return existing + (projected / planning_horizon) * t
        r_rate = np.log((existing + projected) / existing) / planning_horizon
        return existing * np.exp(r_rate * t)
    else:
        return existing + projected

def calculate_ev_adoption(t, existing, projected, planning_horizon, traj_type):
    if traj_type == "Instantaneous":
        return existing + projected
    elif traj_type == "Linear":
        return existing + (projected / planning_horizon) * t
    elif traj_type == "Exponential":
        if existing <= 0:
            return existing + (projected / planning_horizon) * t
        r_rate = np.log((existing + projected) / existing) / planning_horizon
        return existing * np.exp(r_rate * t)
    else:
        return existing + projected

# --- STEP 8: Tariffs Processing ---

def get_tariff_profile(tariff_type, tariff_direction, country_name):
    """
    tariff_type: 'Fixed', 'Dynamic', 'Green Tariff', 'Own Fixed', 'Own Dynamic'
    tariff_direction: 'from' (import/buy), 'to' (export/sell)
    """
    tariff_type = str(tariff_type).strip()
    tariff_direction = str(tariff_direction).strip().lower()
    country_name = str(country_name).strip()
    if tariff_direction not in {"from", "to"}:
        raise ValueError(f"Unknown tariff direction: {tariff_direction!r}")

    # 1. Handle Fixed/Green values (return as list of 35040 same values)
    val = None
    if tariff_type == "Fixed":
        if country_name == "Ukraine":
            if tariff_direction == "from":
                val = tariff_fixed_from_ua
            else: val = tariff_fixed_to_ua
        elif country_name == "Latvia":
            if tariff_direction == "from":
                val = tariff_fixed_from_lv
            else: val = tariff_fixed_to_lv
    elif tariff_type == "Green Tariff" and country_name == "Ukraine" and tariff_direction == "to":
        val = tariff_green_ua
    elif tariff_type == "Own Fixed":
        if tariff_direction == "from":
            val = tariff_fixed_from_own
        else: val = tariff_fixed_to_own

    if val is not None:
        if not np.isfinite(float(val)):
            raise ValueError(f"{tariff_type} {tariff_direction} tariff must be finite")
        return [float(val)] * 35040

    # 2. Handle Dynamic profiles (Load from Excel)
    path_to_tariff_file = ""
    target_tariff_file = ""

    if tariff_type == "Dynamic":
        path_to_tariff_file = os.path.join("Database", "Tariffs")
        prefix = "ua" if country_name == "Ukraine" else "lv"
        target_tariff_file = f"{prefix}_dynamic_{tariff_direction}.xlsx"
    elif tariff_type == "Own Dynamic":
        path_to_tariff_file = os.path.join("User Own Input", "Tariffs")
        target_tariff_file = f"own_dynamic_{tariff_direction}.xlsx"

    else:
        raise ValueError(
            f"Unsupported tariff mode {tariff_type!r} for {tariff_direction} grid direction"
        )

    path = os.path.join(path_to_tariff_file, target_tariff_file)
    return read_tariff_profile(path)

# Execute Tariff Logic
l_tariff_from = get_tariff_profile(tariff_from, "from", country)
l_tariff_to = get_tariff_profile(tariff_to, "to", country)

# --- STEP 9: Output File ---

# 1. Create initial output dictionary with BASE PV and EV data (Year 1)
d_final_output = {
    0: ["Other Input Data", area_type, pv_p_existing, pv_p_projected,
        bess_capacity_existing_kWh, bess_capacity_projected_kWh, pv_adoption_traj],
    1: [header_1, pv_prof] + l_pv_data,
    2: ["PV Aggregated Baseline", f"Factor: {pv_p_existing_scaled:.1f}"] + l_pv_data_scaled,
    3: [header_2, ev_home_share] + d_ev_profiles_results['home'],
    4: [header_3, ev_public_share] + d_ev_profiles_results['public'],
    5: [header_4, ev_work_share] + d_ev_profiles_results['work'],
    6: [header_5, ev_fast_share] + d_ev_profiles_results['fast'],
    7: [header_6, ev_smart_share] + d_ev_profiles_results['smart_home'],
    8: [header_7, ev_smart_share] + d_ev_profiles_results['smart_public'],
    9: ["EV Aggregated Baseline", f"Existing EVs: {ev_num_existing}"] + l_ev_data_aggregated,
    10: ["From Grid, EUR/kWh", tariff_from] + l_tariff_from,
    11: ["To Grid, EUR/kWh ", tariff_to] + l_tariff_to,
    12: ["Run Metadata", json.dumps({
        "country": country,
        "from_grid_tariff_source": tariff_from,
        "to_grid_tariff_source": tariff_to,
        "to_grid_own_fixed_EUR_per_kWh": tariff_fixed_to_own,
        "bess_input_unit": "kWh",
        "bess_degradation_rate": bess_degradation_rate,
    }, ensure_ascii=False)],
}

# 2. Add ECom individual participants' profiles
next_col_idx = 13
for idx, config in enumerate(l_load_data):
    d_final_output[next_col_idx] = config['load_info_header'] + config['load_info_data']
    next_col_idx += 1

# 3. Add aggregated load profiles of participants (Year 1)
d_final_output[next_col_idx] = ["Aggregated Participants Profile", f"{len(l_load_data)} participants"] + l_load_data_aggregated_base
next_col_idx += 1

# To calculate the degradation of PV, create initial list of PV cohorts
# Each cohort: {'capacity': installed capacity, 'install_year': installation year}
# For existing PV, installation year = -pv_age (relative to year 0)
pv_cohorts = [{'capacity': pv_p_existing, 'install_year': -pv_age}]
previous_year_pv_p_installed = pv_p_existing
# print(r"Test: previous_year_pv_p_installed = ", previous_year_pv_p_installed)

# 4. Add aggregated columns for each year of the planning horizon (Dynamic Trajectories)
for year in range(1, planning_horizon + 1):
    year_tag = f"Year {year}"

    # A) Calculate the PV adoption at the beginning of a year
    current_year_pv_p_installed = calculate_pv_adoption(year, pv_p_existing, pv_p_projected, planning_horizon, pv_adoption_traj)
    installed_pv_capacity_t = current_year_pv_p_installed - previous_year_pv_p_installed

    # B) Add a new cohort for the current year
    if installed_pv_capacity_t > 0:
        pv_cohorts.append({'capacity': installed_pv_capacity_t, 'install_year': year})

    previous_year_pv_p_installed = current_year_pv_p_installed
    # print(r"Test: current_year_pv_p_installed = ", previous_year_pv_p_installed)

    # C) Calculate the Baseline vs New effective capacity across all cohorts
    baseline_pv_capacity_t = 0
    new_pv_capacity_t = 0
    for i, cohort in enumerate(pv_cohorts):
        # Cohort's age in the current modeling year
        pv_age_in_year_t = year - cohort['install_year']
        effective_pv_capacity = cohort['capacity'] * (1 - pv_degradation/100) ** pv_age_in_year_t
        if i == 0: # First cohort is baseline
            baseline_pv_capacity_t = effective_pv_capacity
        else: # Subsequent cohorts are new
            new_pv_capacity_t += effective_pv_capacity
        # print(r"Test: cohort in pv_cohorts: cohort['capacity'], PV age in year t, effective_pv_capacity, baseline_pv_capacity_t, new_pv_capacity_t", cohort['capacity'], pv_age_in_year_t, effective_pv_capacity, baseline_pv_capacity_t, new_pv_capacity_t)

    # D) PV Aggregated for year t
    d_final_output[next_col_idx] = [f"PV Aggregated Baseline ({year_tag})", f"Effective kW: {baseline_pv_capacity_t:.2f}"] + [v * baseline_pv_capacity_t for v in l_pv_data]
    next_col_idx += 1
    d_final_output[next_col_idx] = [f"PV Aggregated New ({year_tag})", f"Effective kW: {new_pv_capacity_t:.2f}"] + [v * new_pv_capacity_t for v in l_pv_data]
    next_col_idx += 1

    # E) EV Aggregated for year t
    ev_num_t = calculate_ev_adoption(year, ev_num_existing, ev_num_projected, planning_horizon, ev_adoption_traj)
    l_ev_year_t = []
    for v in l_ev_aggregated_shape:
        value = v * ev_num_t
        l_ev_year_t.append(value)
    d_final_output[next_col_idx] = [f"EV Aggregated ({year_tag})", ev_num_t] + l_ev_year_t
    next_col_idx += 1

    # F) ECom Load Aggregated for year t
    load_growth_factor = (1 + annual_load_growth_rate/100) ** (year - 1)
    l_load_year_t = [v * load_growth_factor for v in l_load_data_aggregated_base]
    d_final_output[next_col_idx] = [f"Total Load ({year_tag})", f"Growth: {load_growth_factor:.2f}"] + l_load_year_t
    next_col_idx += 1

# --- Ensuring that all the arrays are of the same length (to avoid ValueError) ---
target_len = 35042 # 2 lines of headers + 35040
for k in d_final_output:
    current_len = len(d_final_output[k])
    if current_len < target_len:
        d_final_output[k] = d_final_output[k] + [""] * (target_len - current_len)

df_output = pd.DataFrame(d_final_output)


def export_stage1_xlsx(frame, path):
    """Serialize canonical Stage-1 table for legacy MATLAB/download workflows."""
    workbook = openpyxl.Workbook(write_only=True)
    worksheet = workbook.create_sheet(title="Sheet1")
    internal_headers = {"Other Input Data", "Run Metadata"}
    for column_number, header in enumerate(frame.iloc[0], start=1):
        if str(header).strip() in internal_headers:
            column_letter = openpyxl.utils.get_column_letter(column_number)
            worksheet.column_dimensions[column_letter].hidden = True
    for row in frame.itertuples(index=False, name=None):
        worksheet.append([
            None if isinstance(value, (float, np.floating)) and np.isnan(value) else value
            for value in row
        ])
    archive = ZipFile(path, "w", ZIP_DEFLATED, allowZip64=True, compresslevel=1)
    try:
        ExcelWriter(workbook, archive).save()
    finally:
        archive.close()


if os.environ.get("ECOM_STAGE1_SKIP_XLSX") == "1":
    print("[PERF] Stage-1 workbook write: skipped (PreparedData memory path)")
else:
    # Low ZIP compression cuts CPU time for this multi-million-cell workbook.
    write_started = time.perf_counter()
    export_stage1_xlsx(df_output, output_excel)
    print(f"[PERF] Stage-1 workbook write: {time.perf_counter() - write_started:.3f} s")
    print(f"Results successfully saved to {output_excel} for a {planning_horizon}-year horizon. Total columns: {next_col_idx}")
print(f"[PERF] Stage-1 script total: {time.perf_counter() - stage1_started:.3f} s")
